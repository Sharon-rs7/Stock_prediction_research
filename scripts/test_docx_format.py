import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def test():
    doc = docx.Document()
    sec0 = doc.sections[0]
    sec0.top_margin = Inches(0.75)
    sec0.bottom_margin = Inches(0.75)
    sec0.left_margin = Inches(0.75)
    sec0.right_margin = Inches(0.75)

    p = doc.add_paragraph('Machine Learning Framework for Stock Price Forecasting')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add continuous section break for 2 columns
    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    body_sec.top_margin = Inches(0.75)
    body_sec.bottom_margin = Inches(0.75)
    body_sec.left_margin = Inches(0.75)
    body_sec.right_margin = Inches(0.75)
    
    sectPr = body_sec._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), '2')
        cols[0].set(qn('w:space'), '360')
    else:
        cols_xml = f'<w:cols {nsdecls("w")} w:num="2" w:space="360"/>'
        sectPr.append(parse_xml(cols_xml))

    doc.add_paragraph('Column 1 text here.')
    doc.save('test_col.docx')
    print('Test DOCX created successfully!')

if __name__ == '__main__':
    test()
